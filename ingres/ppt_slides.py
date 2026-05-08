import os
from llama_index.readers.file import PptxReader
import logging
from pathlib import Path
from typing import Dict, List, Optional
from fsspec import AbstractFileSystem
from llama_index.core.schema import Document
from pptx import Presentation

log = logging.getLogger(__name__)

class PptxSlideReader(PptxReader):
    
    def __init__(self):
        super(PptxSlideReader, self).__init__()
    
    def load_data(
        self,
        file: Path,
        extra_info: Optional[Dict] = None,
        fs: Optional[AbstractFileSystem] = None,
    ) -> List[Document]:
        """Parse file."""

        if fs:
            with fs.open(file) as f:
                presentation = Presentation(f)
        else:
            presentation = Presentation(file)
        result = ""
        documents = []
        for i, slide in enumerate(presentation.slides):
            result += f"\n\nSlide #{i}: \n"
            for shape in slide.shapes:
                if hasattr(shape, "image"):
                    image = shape.image
                    # get image "file" contents
                    image_bytes = image.blob
                    # temporarily save the image to feed into model
                    image_filename = f"tmp_image.{image.ext}"
                    with open(image_filename, "wb") as f:
                        f.write(image_bytes)
                    result += f"\n Image: {self.caption_image(image_filename)}\n\n"

                    os.remove(image_filename)
                if hasattr(shape, "text"):
                    result += f"{shape.text}\n"
            metadata = extra_info or {}
            metadata["level"] = 1
            metadata["h1"] = "Slide {}".format(i)
            metadata["h2"] = ""
            metadata["h3"] = ""
            metadata["category"] = "PPT"
            document = Document(
                    text=result,
                    metadata=metadata,
                )
            documents.append(document)
            
        return documents